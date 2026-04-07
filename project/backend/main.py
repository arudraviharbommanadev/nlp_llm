import json
import re
from io import BytesIO
from pathlib import Path

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from docx import Document

from backend.chatbot import OllamaChatbot
from backend.database import get_session, init_db, list_sessions, save_session


app = FastAPI(title="Offline Ollama Chatbot API")
chatbot = OllamaChatbot()
BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def on_startup() -> None:
    init_db()


app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")


class ChatRequest(BaseModel):
    message: str = Field(..., min_length=1, description="The user's chat message.")


class ChatResponse(BaseModel):
    response: str
    intent: str
    model: str


class ChatMessage(BaseModel):
    role: str = Field(..., pattern="^(user|assistant)$")
    content: str = Field(..., min_length=1)


class EndSessionRequest(BaseModel):
    messages: list[ChatMessage]


class SessionSummary(BaseModel):
    id: int
    title: str
    created_at: str


class SessionDetail(SessionSummary):
    messages: list[ChatMessage]


def build_transcript_lines(session: dict[str, object]) -> list[str]:
    messages = session["messages"]
    lines = [
        f"Session Title: {session['title']}",
        f"Created At: {session['created_at']}",
        "",
    ]

    for message in messages:
        role = str(message["role"]).capitalize()
        content = str(message["content"])
        lines.extend([f"{role}:", content, ""])

    return lines


def sanitize_filename(title: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", title.strip()).strip("-")
    return cleaned[:64] or "chat-session"


def wrap_text(text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
    words = text.split()
    if not words:
        return [""]

    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if stringWidth(candidate, font_name, font_size) <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def export_session_pdf(session: dict[str, object]) -> BytesIO:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    left_margin = 48
    top_margin = height - 56
    line_height = 16
    max_width = width - (left_margin * 2)
    y_position = top_margin

    pdf.setTitle(str(session["title"]))
    pdf.setFont("Helvetica", 11)

    for raw_line in build_transcript_lines(session):
        wrapped_lines = wrap_text(raw_line, "Helvetica", 11, max_width)
        for line in wrapped_lines:
            if y_position <= 56:
                pdf.showPage()
                pdf.setFont("Helvetica", 11)
                y_position = top_margin

            pdf.drawString(left_margin, y_position, line)
            y_position -= line_height

    pdf.save()
    buffer.seek(0)
    return buffer


def export_session_docx(session: dict[str, object]) -> BytesIO:
    document = Document()
    document.add_heading(str(session["title"]), level=1)
    document.add_paragraph(f"Created At: {session['created_at']}")

    for message in session["messages"]:
        role = str(message["role"]).capitalize()
        document.add_heading(role, level=2)
        document.add_paragraph(str(message["content"]))

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/")
def serve_frontend() -> FileResponse:
    return FileResponse(FRONTEND_DIR / "index.html")


@app.post("/chat", response_model=ChatResponse)
def chat(request: ChatRequest) -> ChatResponse:
    try:
        result = chatbot.process_message(request.message)
        return ChatResponse(
            response=result["response"],
            intent=result["intent"],
            model=result["model"],
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # pragma: no cover
        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to generate a response from Ollama. "
                "Make sure Ollama is installed, running, and the model is available locally."
            ),
        ) from exc


@app.post("/chat/stream")
def stream_chat(request: ChatRequest) -> StreamingResponse:
    def event_stream() -> str:
        try:
            for event in chatbot.stream_response(request.message):
                yield json.dumps(event) + "\n"
        except ValueError as exc:
            yield json.dumps({"type": "error", "detail": str(exc)}) + "\n"
        except Exception:  # pragma: no cover
            yield json.dumps(
                {
                    "type": "error",
                    "detail": (
                        "Failed to generate a response from Ollama. "
                        "Make sure Ollama is installed, running, and the model is available locally."
                    ),
                }
            ) + "\n"

    return StreamingResponse(event_stream(), media_type="application/x-ndjson")


@app.get("/sessions", response_model=list[SessionSummary])
def get_sessions() -> list[SessionSummary]:
    return [SessionSummary(**session) for session in list_sessions()]


@app.get("/sessions/{session_id}", response_model=SessionDetail)
def get_saved_session(session_id: int) -> SessionDetail:
    session = get_session(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    return SessionDetail(**session)


@app.get("/sessions/{session_id}/export")
def export_session(session_id: int, format: str) -> StreamingResponse:
    session = get_session(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")

    normalized_format = format.lower()
    filename_root = sanitize_filename(str(session["title"]))

    if normalized_format == "pdf":
        media_type = "application/pdf"
        buffer = export_session_pdf(session)
        filename = f"{filename_root}.pdf"
    elif normalized_format == "docx":
        media_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        buffer = export_session_docx(session)
        filename = f"{filename_root}.docx"
    else:
        raise HTTPException(status_code=400, detail="Format must be pdf or docx.")

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(buffer, media_type=media_type, headers=headers)


@app.post("/sessions/end", response_model=SessionSummary)
def end_session(request: EndSessionRequest) -> SessionSummary:
    if not request.messages:
        raise HTTPException(status_code=400, detail="Cannot save an empty session.")

    saved = save_session([message.model_dump() for message in request.messages])
    return SessionSummary(**saved)
